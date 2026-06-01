from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_DIR = Path(__file__).resolve().parent.parent
INPUT_CSV = PROJECT_DIR / "data" / "processed" / "analysis_model_data.csv"
RAW_CSV = PROJECT_DIR / "data" / "raw" / "users_dataset.csv"
OUTPUT_XLSX = PROJECT_DIR / "outputs" / "tables" / "appendix_core_sample_50.xlsx"

SAMPLE_SIZE = 50
RANDOM_SEED = 20260601

CHINESE_HEADERS = {
    "sample_no": "样本序号",
    "source_row_no": "原始CSV行号",
    "country": "国家/地区",
    "countryCode": "国家代码",
    "gender": "性别",
    "productsSold": "售出商品数",
    "productsListed": "上架商品数",
    "socialNbFollowers": "粉丝数",
    "socialProductsLiked": "获赞数",
    "seniority": "资历",
    "productsSold_log": "售出商品数对数",
    "productsListed_log": "上架商品数对数",
    "socialNbFollowers_log": "粉丝数对数",
    "socialProductsLiked_log": "获赞数对数",
    "is_female": "女性虚拟变量",
    "seniority_centered": "中心化资历",
    "seniority_centered_sq": "中心化资历平方项",
    "sold_rate": "售出率",
}


def autosize_and_style(workbook):
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True, color="1F1F1F")
    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = border

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border

        for col_idx, col_cells in enumerate(worksheet.columns, start=1):
            values = [str(cell.value) if cell.value is not None else "" for cell in col_cells]
            width = min(max(max(len(value) for value in values) + 2, 10), 42)
            worksheet.column_dimensions[get_column_letter(col_idx)].width = width


def build_sample():
    model_df = pd.read_csv(INPUT_CSV)
    raw_df = pd.read_csv(RAW_CSV, usecols=["countryCode"])

    if len(model_df) != len(raw_df):
        raise ValueError(
            f"Processed rows ({len(model_df)}) and raw rows ({len(raw_df)}) do not match."
        )

    model_df = model_df.copy()
    model_df.insert(0, "source_row_no", np.arange(2, len(model_df) + 2))
    model_df.insert(1, "countryCode", raw_df["countryCode"].values)

    columns = [
        "source_row_no",
        "country",
        "countryCode",
        "gender",
        "productsSold",
        "productsListed",
        "socialNbFollowers",
        "socialProductsLiked",
        "seniority",
        "productsSold_log",
        "productsListed_log",
        "socialNbFollowers_log",
        "socialProductsLiked_log",
        "is_female",
        "seniority_centered",
        "seniority_centered_sq",
        "sold_rate",
    ]

    sample_df = (
        model_df.sample(n=SAMPLE_SIZE, random_state=RANDOM_SEED)
        .loc[:, columns]
        .reset_index(drop=True)
    )
    sample_df.insert(0, "sample_no", np.arange(1, SAMPLE_SIZE + 1))

    round_columns = [
        "productsSold_log",
        "productsListed_log",
        "socialNbFollowers_log",
        "socialProductsLiked_log",
        "seniority_centered",
        "seniority_centered_sq",
        "sold_rate",
    ]
    sample_df[round_columns] = sample_df[round_columns].round(6)
    sample_df = sample_df.rename(columns=CHINESE_HEADERS)
    return sample_df, len(model_df)


def build_notes(total_rows):
    sampling_note = pd.DataFrame(
        [
            ("完整数据来源", "data/raw/users_dataset.csv"),
            ("核心分析数据", "data/processed/analysis_model_data.csv"),
            ("完整样本量", total_rows),
            ("抽样样本量", SAMPLE_SIZE),
            ("随机种子", RANDOM_SEED),
            (
                "抽样方法",
                "使用 pandas.DataFrame.sample(n=50, random_state=20260601) "
                "对核心分析数据进行简单随机抽样。",
            ),
            (
                "处理口径",
                "核心连续变量已按论文分析脚本进行 1%/99% 分位数缩尾，并生成对数、"
                "虚拟变量、中心化项和售出率。",
            ),
            ("生成日期", "2026-06-01"),
        ],
        columns=["项目", "内容"],
    )

    variable_info = pd.DataFrame(
        [
            ("sample_no", "样本序号", "本附录表内的顺序编号。"),
            (
                "source_row_no",
                "原始CSV行号",
                "对应 data/raw/users_dataset.csv 中的行号，含表头行。",
            ),
            ("country", "国家/地区", "用户所在国家或地区。"),
            ("countryCode", "国家代码", "原始数据中的国家代码。"),
            ("gender", "性别", "原始数据中的性别标识，F 表示女性，M 表示男性。"),
            (
                "productsSold",
                "售出商品数",
                "进入模型前经 1% 与 99% 分位数缩尾处理后的售出商品数。",
            ),
            (
                "productsListed",
                "上架商品数",
                "进入模型前经 1% 与 99% 分位数缩尾处理后的上架商品数。",
            ),
            (
                "socialNbFollowers",
                "粉丝数",
                "进入模型前经 1% 与 99% 分位数缩尾处理后的粉丝数。",
            ),
            (
                "socialProductsLiked",
                "获赞数",
                "进入模型前经 1% 与 99% 分位数缩尾处理后的商品获赞数。",
            ),
            (
                "seniority",
                "资历",
                "进入模型前经 1% 与 99% 分位数缩尾处理后的平台注册/使用资历。",
            ),
            ("productsSold_log", "售出商品数对数", "ln(productsSold + 1)。"),
            ("productsListed_log", "上架商品数对数", "ln(productsListed + 1)。"),
            ("socialNbFollowers_log", "粉丝数对数", "ln(socialNbFollowers + 1)。"),
            ("socialProductsLiked_log", "获赞数对数", "ln(socialProductsLiked + 1)。"),
            ("is_female", "女性虚拟变量", "gender 为 F 时取 1，否则取 0。"),
            ("seniority_centered", "中心化资历", "seniority 减去样本均值。"),
            ("seniority_centered_sq", "中心化资历平方项", "seniority_centered 的平方。"),
            (
                "sold_rate",
                "售出率",
                "productsListed 大于 0 时为 productsSold / productsListed；否则为空值。",
            ),
        ],
        columns=["variable", "中文名称", "说明"],
    )
    return sampling_note, variable_info


def main():
    sample_df, total_rows = build_sample()
    sampling_note, variable_info = build_notes(total_rows)

    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        sample_df.to_excel(writer, sheet_name="核心样本50条", index=False)
        sampling_note.to_excel(writer, sheet_name="抽样说明", index=False)
        variable_info.to_excel(writer, sheet_name="变量说明", index=False)
        autosize_and_style(writer.book)

    print(OUTPUT_XLSX)
    print(f"rows={len(sample_df)} seed={RANDOM_SEED}")


if __name__ == "__main__":
    main()
